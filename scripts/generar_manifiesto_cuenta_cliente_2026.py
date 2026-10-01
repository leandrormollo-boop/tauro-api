#!/usr/bin/env python3
"""Genera el manifiesto para completar la cuenta corriente 2026 de un cliente.

Fuente de verdad: la planilla del cliente (lo efectivamente facturado). La
maestra TAURO 2026 sólo se usa para completar una fecha faltante; nunca pisa
un importe. El manifiesto resultante NO se versiona porque contiene datos del
cliente; la aplicación (servicios/importacion_cuenta_cliente.py) vuelve a
validar su huella y sus totales antes de escribir.

Reglas:
  * Cada FLETE/IMPO/REVISAR/CANCELADO es un envío. Los TAX de la planilla se
    suman al envío con el mismo tracking (el portal los muestra en una fila
    aparte). Un TAX sin flete visible queda en un envío oculto "ancla".
  * Las diferencias de la columna DIFERENCIAS se trasladan tal cual.
  * Un tracking repetido en la planilla es un posible doble cobro:
      - MELCIOR: sólo se carga la primera fila; las demás quedan informadas
        para que Leandro decida (no se crean).
      - PRETE ROSSO: se respeta la regla del lote auditado por Codex: la
        repetición se carga como movimiento contable oculto.
  * El saldo pendiente 2025 y los pagos salen de las pestañas de detalle.

Uso:
  python3 scripts/generar_manifiesto_cuenta_cliente_2026.py --cliente MELCIOR \
      --planilla "MELCIOR 2026 TAURO.xlsx" --maestra "TAURO 2026.xlsx" \
      --output manifiesto_melcior_2026.json
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from auditar_melcior_2026 import auditar  # noqa: E402

MESES = (
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
)
CENTAVO = Decimal("0.01")
TIPOS_BASE = {"FLETE", "IMPO", "REVISAR", "CANCELADO"}

CLIENTES = {
    "MELCIOR": {
        "cliente_id": "MELCIOR",
        "prefijo": "MELCIOR-2026",
        "remitente": "JUAN PABLO MELCIOR",
        "duplicados": "NO_CREAR",
    },
    "PRETE ROSSO": {
        "cliente_id": "PRETE ROSSO",
        "prefijo": "PRETE-ROSSO-2026",
        "remitente": "PRETE ROSSO",
        "duplicados": "CARGAR_OCULTO",
    },
}


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return " ".join(str(valor).strip().split())


def _dinero(valor: Any) -> Decimal:
    if valor in (None, ""):
        return Decimal("0.00")
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor)).quantize(CENTAVO)
    bruto = str(valor).replace("$", "").replace(" ", "").replace(" ", "")
    if "," in bruto and "." in bruto:
        bruto = bruto.replace(",", "") if bruto.rfind(".") > bruto.rfind(",") else bruto.replace(".", "").replace(",", ".")
    elif "," in bruto:
        bruto = bruto.replace(",", ".")
    try:
        return Decimal(bruto).quantize(CENTAVO)
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _fecha(valor: Any) -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = _texto(valor)
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            pass
    return ""


def _tracking(valor: Any) -> str:
    texto = _texto(valor).replace(" ", "")
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto if re.fullmatch(r"\d{8,20}", texto) else ""


def _hash_archivo(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _hash_json(datos: Any) -> str:
    return hashlib.sha256(json.dumps(
        datos, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str
    ).encode("utf-8")).hexdigest()


# ── Lectura de planillas ────────────────────────────────────────────────────

def _datos_maestra(path: Path, cliente_tauro: str) -> tuple[dict[str, str], dict[str, str], dict[str, dict]]:
    """tracking -> fecha, courier físico y costo según TAURO 2026.

    La fecha sólo completa una faltante y nunca se pisan importes al cliente.
    BOXFLY y ORION despachan por FedEx: el portal muestra el courier real y
    nunca el nombre del operador. El costo (SALDO o, si falta, COSTOINICIAL)
    es interno: sirve para la rentabilidad y el cliente no lo ve.
    """
    ws = load_workbook(path, data_only=True, read_only=True)["ENVIOS 2026"]
    fechas: dict[str, str] = {}
    couriers: dict[str, str] = {}
    costos: dict[str, dict] = defaultdict(lambda: {"flete": [], "tax": Decimal("0.00"), "tax_sin_costo": 0})
    for fila in ws.iter_rows(min_row=2, max_col=16, values_only=True):
        if _texto(fila[4]).upper() != cliente_tauro:
            continue
        tracking = _tracking(fila[9])
        if not tracking:
            continue
        fecha = _fecha(fila[3])
        if fecha and tracking not in fechas:
            fechas[tracking] = fecha
        empresa = _texto(fila[0]).upper()
        courier = "DHL" if "DHL" in empresa else ("UPS" if "UPS" in empresa else "FEDEX")
        couriers.setdefault(tracking, courier)
        tipo = _texto(fila[10]).upper() or "FLETE"
        costo = _dinero(fila[14]) or _dinero(fila[15])
        if tipo == "TAX":
            if costo > 0:
                costos[tracking]["tax"] += costo
            else:
                costos[tracking]["tax_sin_costo"] += 1
        elif tipo == "FLETE":
            costos[tracking]["flete"].append(costo)
    return fechas, couriers, dict(costos)


def _costo_envio(costos: dict[str, dict], tracking: str, tax_cliente: Decimal, *, solo_tax: bool = False) -> str:
    """Costo total (flete + TAX) según TAURO, o "" si no es inequívoco."""
    datos = costos.get(tracking)
    if not datos:
        return ""
    total = Decimal("0.00")
    if not solo_tax:
        if len(datos["flete"]) != 1 or datos["flete"][0] <= 0:
            return ""
        total += datos["flete"][0]
    if tax_cliente > 0:
        if datos["tax_sin_costo"] or datos["tax"] <= 0:
            return ""
        total += datos["tax"]
    return str(total.quantize(CENTAVO)) if total > 0 else ""


def _leer_melcior(path: Path) -> tuple[list[dict], list[dict], Decimal, Decimal | None]:
    auditoria = auditar(path)
    filas = []
    for r in auditoria["filas"]:
        tipo = (r.get("tipo") or "FLETE").upper()
        tipo = "CANCELADO" if tipo == "CANCELADA" else tipo
        filas.append({
            "mes": r["mes"], "fila_cliente": int(r["fila"]), "oculta": False,
            "fecha": r.get("fecha") or "", "remitente": r.get("remitente") or "",
            "destinatario": r.get("destinatario") or "", "pais_fuente": r.get("pais") or "",
            "peso_fuente": r.get("peso") or "", "medidas_fuente": r.get("medidas") or "",
            "tracking": _tracking(r.get("tracking")), "tipo_fuente": tipo,
            "importe_ars": _dinero(r.get("facturado")),
            "diferencia_flete_ars": _dinero(r.get("diferencia")),
            "nro_cliente_fuente": r.get("fc") or "",
        })
    detalle = auditoria.get("detalle_2026") or {}
    pagos = [{
        "source_key": f"MELCIOR-2026:PAGO:{p['fila']}",
        "fila_cliente": int(p["fila"]),
        # La planilla no informa fechas de pago. La importación base usó el
        # 02/09/2026 marcando la fecha como no documentada; se conserva.
        "fecha": "2026-09-02",
        "monto_ars": str(_dinero(p["monto_ars"])),
        "referencia": "",
        "detalle_fuente": p.get("detalle") or "",
        "fecha_original_informada": False,
    } for p in detalle.get("pagos", []) if _dinero(p["monto_ars"]) > 0]
    saldo_2025 = _dinero(detalle.get("saldo_pendiente_2025"))
    wb = load_workbook(path, data_only=True, read_only=True)
    saldo_planilla = _dinero(wb["DETALLE 2026"]["A6"].value) if "DETALLE 2026" in wb.sheetnames else None
    return filas, pagos, saldo_2025, saldo_planilla


def _config_prete(mes: str) -> dict[str, int | None]:
    if mes in {"ENERO", "FEBRERO", "MARZO"}:
        return {"fecha": 1, "remitente": 2, "destinatario": 3, "pais": 4, "peso": 5,
                "medidas": 6, "tipo": 7, "tracking": 8, "facturado": 10,
                "diferencia": None, "nro_cliente": 11}
    return {"fecha": 2, "remitente": 3, "destinatario": 4, "pais": 5, "peso": 6,
            "medidas": 7, "tracking": 8, "tipo": 9, "facturado": 10,
            "diferencia": 11, "nro_cliente": 12}


def _leer_prete(path: Path) -> tuple[list[dict], list[dict], Decimal, Decimal | None]:
    wb = load_workbook(path, data_only=True)
    filas = []
    for mes in MESES:
        if mes not in wb.sheetnames:
            continue
        ws = wb[mes]
        cfg = _config_prete(mes)
        for numero in range(6, ws.max_row + 1):
            tipo = _texto(ws.cell(numero, int(cfg["tipo"])).value).upper()
            tipo = "CANCELADO" if tipo == "CANCELADA" else tipo
            tracking = _tracking(ws.cell(numero, int(cfg["tracking"])).value)
            if not tracking or tipo not in TIPOS_BASE | {"TAX"}:
                continue
            importe_fuente = ws.cell(numero, int(cfg["facturado"])).value
            # Un CANCELADO nunca genera deuda (en AGOSTO hay texto accidental
            # dentro del importe de un cancelado).
            importe = Decimal("0.00") if tipo == "CANCELADO" else _dinero(importe_fuente)
            col_dif = cfg["diferencia"]
            filas.append({
                "mes": mes, "fila_cliente": numero,
                "oculta": bool(ws.row_dimensions[numero].hidden),
                "fecha": _fecha(ws.cell(numero, int(cfg["fecha"])).value),
                "remitente": _texto(ws.cell(numero, int(cfg["remitente"])).value),
                "destinatario": _texto(ws.cell(numero, int(cfg["destinatario"])).value),
                "pais_fuente": _texto(ws.cell(numero, int(cfg["pais"])).value),
                "peso_fuente": _texto(ws.cell(numero, int(cfg["peso"])).value),
                "medidas_fuente": _texto(ws.cell(numero, int(cfg["medidas"])).value),
                "tracking": tracking, "tipo_fuente": tipo, "importe_ars": importe,
                "diferencia_flete_ars": _dinero(ws.cell(numero, int(col_dif)).value) if col_dif else Decimal("0.00"),
                "nro_cliente_fuente": _texto(ws.cell(numero, int(cfg["nro_cliente"])).value),
            })
    pagos = []
    detalle = wb["DETALLE 2026"]
    for numero in range(2, detalle.max_row + 1):
        monto = _dinero(detalle.cell(numero, 1).value)
        valor_fecha = detalle.cell(numero, 3).value
        if isinstance(valor_fecha, (int, float)) and 30000 < float(valor_fecha) < 80000:
            from openpyxl.utils.datetime import from_excel
            valor_fecha = from_excel(float(valor_fecha))
        fecha = _fecha(valor_fecha)
        if monto <= 0 or not fecha:
            continue
        pagos.append({
            "source_key": f"PRETE-ROSSO-2026:PAGO:{numero}",
            "fila_cliente": numero, "fecha": fecha, "monto_ars": str(monto),
            "referencia": _texto(detalle.cell(numero, 2).value),
            "detalle_fuente": "", "fecha_original_informada": True,
        })
    saldo_2025 = _dinero(wb["pendiente 2025"]["B4"].value) if "pendiente 2025" in wb.sheetnames else Decimal("0.00")
    return filas, pagos, saldo_2025, None


# ── Armado del manifiesto ───────────────────────────────────────────────────

def generar(cliente: str, planilla: Path, maestra: Path, *, incluir_ocultas: bool = False) -> dict[str, Any]:
    cfg = CLIENTES[cliente]
    lector = _leer_melcior if cliente == "MELCIOR" else _leer_prete
    filas, pagos, saldo_2025, saldo_planilla = lector(planilla)
    fechas_tauro, couriers, costos = _datos_maestra(maestra, "JUAN PABLO MELCIOR" if cliente == "MELCIOR" else "PRETE ROSSO")
    excluidas: list[dict] = []
    vigentes = []
    for fila in filas:
        if fila["oculta"] and not incluir_ocultas:
            excluidas.append({"fila": f"{fila['mes']}!{fila['fila_cliente']}", "motivo": "FILA_OCULTA"})
            continue
        if not fila["fecha"] and fila["tracking"] in fechas_tauro:
            fila["fecha"] = fechas_tauro[fila["tracking"]]
            fila["fecha_completada_desde"] = "TAURO 2026"
        vigentes.append(fila)

    tax_por_tracking: dict[str, list[dict]] = defaultdict(list)
    bases = []
    for fila in vigentes:
        if fila["tipo_fuente"] == "TAX":
            if fila["importe_ars"] > 0:
                tax_por_tracking[fila["tracking"]].append(fila)
            continue
        if fila["tipo_fuente"] in TIPOS_BASE:
            bases.append(fila)

    repeticiones = defaultdict(int)
    for fila in bases:
        if fila["tipo_fuente"] != "CANCELADO":
            repeticiones[fila["tracking"]] += 1
    ocurrencias = defaultdict(int)
    envios: list[dict[str, Any]] = []
    for base in bases:
        cancelado = base["tipo_fuente"] == "CANCELADO"
        ocurrencias[base["tracking"]] += 1
        n = ocurrencias[base["tracking"]]
        es_repetido = not cancelado and repeticiones[base["tracking"]] > 1 and n > 1
        if cliente == "MELCIOR":
            source_key = f"{cfg['prefijo']}:{base['mes']}:{base['fila_cliente']}"
        else:
            source_key = f"{cfg['prefijo']}:{base['mes']}:{base['fila_cliente']}:{n}"
        tax = Decimal("0.00")
        tax_fuentes: list[str] = []
        if not cancelado and not es_repetido and base["tracking"] in tax_por_tracking:
            impuestos = tax_por_tracking.pop(base["tracking"])
            tax = sum((f["importe_ars"] for f in impuestos), Decimal("0.00"))
            tax_fuentes = [f"{f['mes']}!{f['fila_cliente']}" for f in impuestos]
        importe = Decimal("0.00") if cancelado else base["importe_ars"]
        diferencia = Decimal("0.00") if cancelado else base["diferencia_flete_ars"]
        crear = True
        visible = True
        if es_repetido:
            visible = False
            crear = cfg["duplicados"] == "CARGAR_OCULTO"
        envios.append({
            "source_key": source_key,
            "mes": base["mes"],
            "fila_cliente": base["fila_cliente"],
            "fecha": base["fecha"],
            "fecha_completada_desde": base.get("fecha_completada_desde", ""),
            "remitente": base["remitente"] or cfg["remitente"],
            "destinatario": base["destinatario"] or "Destinatario no informado",
            "pais_fuente": base["pais_fuente"],
            "peso_fuente": base["peso_fuente"],
            "medidas_fuente": base["medidas_fuente"],
            "tracking": base["tracking"],
            "courier": couriers.get(base["tracking"], "FEDEX"),
            "tipo_fuente": base["tipo_fuente"],
            "estado_portal": "CANCELADO" if cancelado else "DESPACHADO",
            "visible_cliente": visible,
            "tracking_repetido": es_repetido,
            "crear_si_falta": crear,
            "genera_deuda": bool(importe > 0 and not cancelado),
            "requiere_revision": bool(es_repetido or base["tipo_fuente"] == "REVISAR" or (not cancelado and importe == 0)),
            "importe_inicial_ars": str(importe),
            "diferencia_flete_ars": str(diferencia),
            "tax_cliente_ars": str(tax),
            "tax_fuentes": tax_fuentes,
            "nro_cliente_fuente": base["nro_cliente_fuente"],
            "costo_courier_ars": "" if (cancelado or es_repetido) else _costo_envio(costos, base["tracking"], tax),
        })

    # TAX sin flete visible: ancla contable oculta (el TAX sí aparece en la cuenta).
    for tracking, impuestos in sorted(tax_por_tracking.items()):
        primera = impuestos[0]
        envios.append({
            "source_key": f"{cfg['prefijo']}:{primera['mes']}:TAX:{primera['fila_cliente']}",
            "mes": primera["mes"], "fila_cliente": primera["fila_cliente"],
            "fecha": primera["fecha"], "fecha_completada_desde": primera.get("fecha_completada_desde", ""),
            "remitente": primera["remitente"] or cfg["remitente"],
            "destinatario": primera["destinatario"] or "TAX histórico",
            "pais_fuente": primera["pais_fuente"], "peso_fuente": primera["peso_fuente"],
            "medidas_fuente": primera["medidas_fuente"], "tracking": tracking,
            "courier": couriers.get(tracking, "FEDEX"),
            "tipo_fuente": "TAX_SIN_FLETE", "estado_portal": "DESPACHADO",
            "visible_cliente": False, "tracking_repetido": False, "crear_si_falta": True,
            "genera_deuda": False, "requiere_revision": True,
            "importe_inicial_ars": "0.00", "diferencia_flete_ars": "0.00",
            "tax_cliente_ars": str(sum((f["importe_ars"] for f in impuestos), Decimal("0.00"))),
            "tax_fuentes": [f"{f['mes']}!{f['fila_cliente']}" for f in impuestos],
            "nro_cliente_fuente": "",
            "costo_courier_ars": _costo_envio(
                costos, tracking, sum((f["importe_ars"] for f in impuestos), Decimal("0.00")), solo_tax=True),
        })

    envios.sort(key=lambda f: (MESES.index(f["mes"]), f["fila_cliente"], f["source_key"]))
    sin_fecha = [f["source_key"] for f in envios if not f["fecha"]]
    if sin_fecha:
        raise ValueError(f"Envíos sin fecha (completar en la planilla): {sin_fecha[:10]}")

    def total(filas_, campo, solo_creables=True):
        return sum((Decimal(f[campo]) for f in filas_ if f["crear_si_falta"] or not solo_creables), Decimal("0.00"))

    resumen_mensual = {}
    for mes in MESES:
        filas_mes = [f for f in envios if f["mes"] == mes]
        if not filas_mes:
            continue
        resumen_mensual[mes] = {
            "envios": len(filas_mes),
            "cargos": sum(f["genera_deuda"] and f["crear_si_falta"] for f in filas_mes),
            "cancelados": sum(f["estado_portal"] == "CANCELADO" for f in filas_mes),
            "repetidos_sin_crear": sum(not f["crear_si_falta"] for f in filas_mes),
            "fletes_ars": str(total(filas_mes, "importe_inicial_ars")),
            "diferencias_ars": str(total(filas_mes, "diferencia_flete_ars")),
            "tax_ars": str(total(filas_mes, "tax_cliente_ars")),
        }
    pagos_total = sum((Decimal(p["monto_ars"]) for p in pagos), Decimal("0.00"))
    fletes = total(envios, "importe_inicial_ars")
    difs = total(envios, "diferencia_flete_ars")
    taxs = total(envios, "tax_cliente_ars")
    repetidos = [f for f in envios if not f["crear_si_falta"]]
    contenido: dict[str, Any] = {
        "schema_version": 2,
        "tipo": "CUENTA_CLIENTE_2026",
        "cliente_id": cfg["cliente_id"],
        "periodo": 2026,
        "generado_at": datetime.now().replace(microsecond=0).isoformat(),
        "source_files_sha256": {
            "planilla_cliente": _hash_archivo(planilla),
            "tauro_2026": _hash_archivo(maestra),
        },
        "politica_duplicados": cfg["duplicados"],
        "incluye_filas_ocultas": incluir_ocultas,
        "filas_excluidas": excluidas,
        "envios": envios,
        "saldo_pendiente_2025": {
            "source_key": f"{cfg['prefijo']}:SALDO-PENDIENTE-2025",
            "fecha": "2025-12-31",
            "monto_ars": str(saldo_2025),
            "concepto": "SALDO PENDIENTE 2025",
        },
        "pagos": pagos,
        "resumen_mensual": resumen_mensual,
        "resumen": {
            "envios": len(envios),
            "cargos": sum(f["genera_deuda"] and f["crear_si_falta"] for f in envios),
            "cancelados": sum(f["estado_portal"] == "CANCELADO" for f in envios),
            "repetidos_sin_crear": len(repetidos),
            "repetidos_sin_crear_ars": str(sum((
                Decimal(f["importe_inicial_ars"]) + Decimal(f["diferencia_flete_ars"]) + Decimal(f["tax_cliente_ars"])
                for f in repetidos), Decimal("0.00"))),
            "fletes_ars": str(fletes),
            "diferencias_ars": str(difs),
            "tax_ars": str(taxs),
            "saldo_pendiente_2025_ars": str(saldo_2025),
            "pagos": len(pagos),
            "pagos_ars": str(pagos_total),
            "saldo_resultante_ars": str(fletes + difs + taxs + saldo_2025 - pagos_total),
            "saldo_segun_planilla_ars": str(saldo_planilla) if saldo_planilla is not None else "",
        },
    }
    contenido["manifest_sha256"] = _hash_json(contenido)
    return contenido


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cliente", required=True, choices=sorted(CLIENTES))
    parser.add_argument("--planilla", type=Path, required=True)
    parser.add_argument("--maestra", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--incluir-ocultas", action="store_true",
                        help="Incluir filas ocultas de la planilla (por defecto se excluyen).")
    args = parser.parse_args()
    manifiesto = generar(args.cliente, args.planilla, args.maestra, incluir_ocultas=args.incluir_ocultas)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(manifiesto, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({
        "manifest_sha256": manifiesto["manifest_sha256"],
        "resumen": manifiesto["resumen"],
        "resumen_mensual": manifiesto["resumen_mensual"],
        "filas_excluidas": len(manifiesto["filas_excluidas"]),
    }, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
