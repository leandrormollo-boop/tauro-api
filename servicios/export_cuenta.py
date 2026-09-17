"""Descarga de los mismos movimientos y filtros que ve el cliente."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from servicios.cuenta_corriente import movimientos_cuenta_paginados


def _celda(hoja, valor, *, encabezado=False):
    celda = WriteOnlyCell(hoja, value=valor)
    # Excel no debe interpretar referencias/guías del usuario como fórmulas.
    if isinstance(valor, str):
        celda.data_type = "s"
    if encabezado:
        celda.font = Font(name="Calibri", bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="352159")
    else:
        celda.font = Font(name="Calibri", size=11)
    celda.alignment = Alignment(vertical="top", wrap_text=True)
    return celda


def generar_excel_cuenta(cliente: str, ambito="consolidado", tipo="todos", **filtros) -> bytes:
    resultado = movimientos_cuenta_paginados(
        cliente, ambito, tipo, exportar=True, **filtros
    )
    libro = Workbook(write_only=True)
    informacion = libro.create_sheet("Consulta")
    informacion.column_dimensions["A"].width = 28
    informacion.column_dimensions["B"].width = 78
    for etiqueta, valor in (
        ("TAURO SOLUTIONS", "Cuenta corriente · importes en ARS"),
        ("Cliente", cliente), ("Ámbito", ambito), ("Tipo", tipo),
        ("Búsqueda", filtros.get("q") or "Todas"),
        ("Desde", filtros.get("desde") or "Sin límite"),
        ("Hasta", filtros.get("hasta") or "Sin límite"),
        ("Movimientos", resultado["total_resultados"]),
        ("Criterio", "Los pagos en revisión, rechazados y envíos cancelados o reemplazados no modifican el saldo."),
        ("Alcance", "Movimientos del período seleccionado; no es un saldo de apertura o cierre."),
    ):
        informacion.append([_celda(informacion, etiqueta), _celda(informacion, valor)])
    hoja = libro.create_sheet("Movimientos")
    hoja.freeze_panes = "A2"
    titulos = (
        "Fecha", "Tipo", "Detalle", "Referencia", "Guía / tracking", "Factura",
        "Ámbito", "Cargo ARS", "Pago / crédito ARS", "Importe informado ARS", "Estado",
    )
    anchos = (15, 24, 42, 25, 24, 26, 21, 21, 23, 25, 20)
    for numero, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(numero)].width = ancho
    hoja.append([_celda(hoja, titulo, encabezado=True) for titulo in titulos])
    for m in resultado["items"]:
        fecha = date.fromisoformat(m["fecha_iso"]) if m.get("fecha_iso") else m.get("fecha", "")
        valores = [
            fecha, m.get("tipo", ""), m.get("etiqueta_envio") or m.get("concepto", ""),
            m.get("referencia") or "", m.get("numero_guia") or "", m.get("numero_factura") or "",
            m.get("ambito") or "", m.get("debe_ars", 0), m.get("haber_ars", 0),
            m.get("monto_ars", 0), m.get("estado") or "",
        ]
        celdas = [_celda(hoja, valor) for valor in valores]
        celdas[0].number_format = "dd/mm/yyyy"
        for i in (7, 8, 9):
            celdas[i].number_format = '#,##0.00;[Red]-#,##0.00'
        hoja.append(celdas)
    hoja.auto_filter.ref = f"A1:K{len(resultado['items']) + 1}"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
