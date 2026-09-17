"""Contrato de presentación de movimientos sobre PostgreSQL real y aislado."""

from __future__ import annotations

import os
import uuid
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path

import psycopg2
import psycopg2.extras
import pytest

from servicios import cuenta_corriente


DATABASE_URL = os.getenv("TAURO_TEST_DATABASE_URL", "").strip()
pytestmark = pytest.mark.skipif(
    not DATABASE_URL,
    reason="requiere TAURO_TEST_DATABASE_URL aislada",
)


@pytest.fixture
def cuenta_db(monkeypatch):
    schema = f"test_cuenta_presentacion_{uuid.uuid4().hex}"
    schema_sql = (
        Path(__file__).resolve().parents[1] / "sql" / "schema.sql"
    ).read_text(encoding="utf-8")
    admin = psycopg2.connect(
        DATABASE_URL,
        cursor_factory=psycopg2.extras.RealDictCursor,
    )
    admin.autocommit = True
    try:
        with admin.cursor() as cur:
            cur.execute(f'CREATE SCHEMA "{schema}"')
            cur.execute(f'SET search_path TO "{schema}"')
            cur.execute(schema_sql)

        @contextmanager
        def get_conn_aislada():
            conn = psycopg2.connect(
                DATABASE_URL,
                cursor_factory=psycopg2.extras.RealDictCursor,
            )
            try:
                with conn.cursor() as cur:
                    cur.execute(f'SET search_path TO "{schema}"')
                yield conn
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()

        monkeypatch.setattr(cuenta_corriente, "get_conn", get_conn_aislada)
        yield get_conn_aislada
    finally:
        with admin.cursor() as cur:
            cur.execute(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
        admin.close()


def test_envio_se_presenta_ordenado_sin_cambiar_el_movimiento(cuenta_db):
    with cuenta_db() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO clientes (cliente_id, email, nombre)
                VALUES ('WAIMAO', 'qa-waimao@example.invalid', 'WAIMAO')
                """
            )
            cur.execute(
                """
                INSERT INTO solicitudes_guia (
                    cliente_id, producto_alias, remitente_nombre,
                    destino_pais, dest_nombre, dest_direccion, dest_ciudad,
                    dest_zip, tracking, ambito
                ) VALUES (
                    'WAIMAO', 'Ropa', 'SLINGER', 'UY', 'MARSANTEX',
                    'Calle QA 123', 'Montevideo', '11000',
                    '6781215324', 'INTERNACIONAL'
                )
                RETURNING id
                """
            )
            solicitud_id = int(cur.fetchone()["id"])
            cur.execute(
                """
                INSERT INTO envios (
                    cliente_id, fecha, monto_ars, estado, descripcion,
                    tracking, solicitud_id, ambito
                ) VALUES (
                    'WAIMAO', DATE '2026-09-02', 1714134, 'ACTIVO',
                    'Flete internacional', '6781215324', %s, 'INTERNACIONAL'
                )
                """,
                (solicitud_id,),
            )

    movimientos = cuenta_corriente.movimientos_cuenta_paginados(
        "WAIMAO", "internacional", "cargos", 1, 10,
    )

    assert movimientos["total_resultados"] == 1
    assert len(movimientos["items"]) == 1
    movimiento = movimientos["items"][0]
    assert movimiento["concepto"] == "Flete"
    assert movimiento["numero_guia"] == "6781215324"
    assert movimiento["destinatario"] == "MARSANTEX"
    assert movimiento["fecha"] == "02/09/2026"
    assert movimiento["remitente"] == "SLINGER"
    assert movimiento["valor_envio_ars"] == Decimal("1714134.00")
    assert movimiento["debe_ars"] == Decimal("1714134.00")
    assert movimiento["haber_ars"] == Decimal("0.00")


def test_busqueda_literal_periodo_y_export_solo_cuenta_propia(cuenta_db):
    from io import BytesIO
    from openpyxl import load_workbook
    from servicios.export_cuenta import generar_excel_cuenta

    with cuenta_db() as conn:
        with conn.cursor() as cur:
            for cliente in ("DUENO", "AJENO"):
                cur.execute("INSERT INTO clientes (cliente_id, nombre, email) VALUES (%s,%s,%s)", (cliente, cliente, cliente.lower() + "@example.invalid"))
            for cliente, fecha, detalle, monto in (
                ("DUENO", "2026-09-02", "Ropa_% & verano", 125),
                ("DUENO", "2026-09-02", "RopaABC", 250),
                ("DUENO", "2026-08-31", "Ropa_% & verano", 375),
                ("AJENO", "2026-09-02", "Ropa_% & verano", 900),
            ):
                cur.execute("""INSERT INTO envios
                    (cliente_id,fecha,monto_ars,estado,descripcion,ambito)
                    VALUES (%s,%s,%s,'ACTIVO',%s,'INTERNACIONAL')""",
                    (cliente, fecha, monto, detalle))
    filtros = {"q": "Ropa_%", "desde": "2026-09-01", "hasta": "2026-09-02"}
    resultado = cuenta_corriente.movimientos_cuenta_paginados("DUENO", **filtros)
    assert resultado["total_resultados"] == 1
    assert resultado["items"][0]["debe_ars"] == Decimal("125.00")
    assert resultado["items"][0]["fecha_iso"] == "2026-09-02"
    libro = load_workbook(BytesIO(generar_excel_cuenta("DUENO", **filtros)))
    filas = list(libro["Movimientos"].values)
    assert len(filas) == 2
    assert filas[1][7] == 125


def test_pagos_rechazados_y_en_revision_visibles_sin_impacto(cuenta_db):
    with cuenta_db() as conn:
        with conn.cursor() as cur:
            cur.execute("INSERT INTO clientes (cliente_id,nombre,email) VALUES ('PAGOS_QA','Prueba','pagos-qa@example.invalid')")
            for estado in ("PENDIENTE", "RECHAZADO", "APROBADO"):
                cur.execute("""INSERT INTO pagos
                    (cliente_id,fecha,monto_ars,metodo,referencia,estado)
                    VALUES ('PAGOS_QA','2026-09-15',100,'Transferencia',%s,%s)""",
                    (estado, estado))
    resultado = cuenta_corriente.movimientos_cuenta_paginados("PAGOS_QA", tipo="pagos")
    assert {m["tipo"] for m in resultado["items"]} == {"PAGO", "PAGO_PENDIENTE", "PAGO_RECHAZADO"}
    no_acreditados = [m for m in resultado["items"] if m["tipo"] != "PAGO"]
    assert all(m["haber_ars"] == m["debe_ars"] == 0 for m in no_acreditados)
    revision = cuenta_corriente.movimientos_cuenta_paginados("PAGOS_QA", tipo="revision")
    assert len(revision["items"]) == 1
    assert revision["items"][0]["estado"] == "PENDIENTE"


def test_cancelados_son_historia_sin_cargo_y_no_ocultan_cargos_activos(cuenta_db):
    from io import BytesIO
    from openpyxl import load_workbook
    from servicios.export_cuenta import generar_excel_cuenta

    with cuenta_db() as conn:
        with conn.cursor() as cur:
            for cliente in ('PROPIO', 'AJENO'):
                cur.execute("INSERT INTO clientes (cliente_id,nombre,email) VALUES (%s,%s,%s)",
                            (cliente, cliente, cliente + '@example.invalid'))
            escenarios = (
                # cliente, estado operativo, estado contable, visible, test
                ('PROPIO', 'CANCELADO', 'CANCELADO', True, False),
                ('PROPIO', 'CANCELADO', None, True, False),
                ('PROPIO', 'REEMPLAZADO', 'CANCELADO', True, False),
                ('PROPIO', 'CANCELADO', 'ACTIVO', True, False),
                ('PROPIO', 'CANCELADO', None, False, False),
                ('PROPIO', 'CANCELADO', None, True, True),
                ('AJENO', 'CANCELADO', None, True, False),
            )
            ids = []
            for i, (cliente, estado, cargo, visible, test) in enumerate(escenarios):
                cur.execute("""INSERT INTO solicitudes_guia
                    (cliente_id,estado,producto_alias,destino_pais,dest_nombre,
                     dest_direccion,dest_ciudad,dest_zip,remitente_ciudad,remitente_pais,
                     tracking,ambito,created_at,visible_cliente,test,numero_guia_tauro)
                    VALUES (%s,%s,'Muestras','AR',%s,'QA 123','Buenos Aires','1000',
                            'Shenzhen','CN',%s,'INTERNACIONAL','2026-09-16 12:00Z',%s,%s,%s)
                    RETURNING id""", (cliente, estado, f'Destino {i}', f'QA-{i}', visible, test, 50300+i))
                sid = cur.fetchone()['id']; ids.append(sid)
                if cargo:
                    cur.execute("""INSERT INTO envios
                        (cliente_id,solicitud_id,fecha,monto_ars,estado,ambito)
                        VALUES (%s,%s,'2026-09-16',900,%s,'INTERNACIONAL')""", (cliente,sid,cargo))
            cur.execute("SELECT id,estado,monto_ars FROM envios ORDER BY id")
            cargos_antes = cur.fetchall()
    filtros = dict(desde='2026-09-01', hasta='2026-09-30')
    todos = cuenta_corriente.movimientos_cuenta_paginados('PROPIO', **filtros)
    assert todos['total_resultados'] == 1
    assert {m['solicitud_id'] for m in todos['items']} == {ids[3]}
    assert sum(m['debe_ars']-m['haber_ars'] for m in todos['items']) == Decimal('900.00')
    bajas = cuenta_corriente.movimientos_cuenta_paginados('PROPIO', tipo='cancelados', **filtros)
    assert bajas['total_resultados'] == 2
    assert {m['solicitud_id'] for m in bajas['items']} == set(ids[:2])
    assert {m['tipo'] for m in bajas['items']} == {'ENVIO_CANCELADO'}
    modificados = cuenta_corriente.movimientos_cuenta_paginados('PROPIO', tipo='modificados', **filtros)
    assert {m['solicitud_id'] for m in modificados['items']} == {ids[2]}
    assert modificados['items'][0]['debe_ars'] == modificados['items'][0]['haber_ars'] == 0
    for m in bajas['items']:
        assert m['debe_ars'] == m['haber_ars'] == m['monto_ars'] == 0
        assert m['valor_envio_ars'] is None
        assert m['origen_ciudad'] == 'Shenzhen'
        assert m['destino_ciudad'] == 'Buenos Aires'
    costos = cuenta_corriente.movimientos_cuenta_paginados('PROPIO', tipo='costos', **filtros)
    assert costos['total_resultados'] == 1
    assert costos['items'][0]['debe_ars'] == Decimal('900.00')
    encontrada = cuenta_corriente.movimientos_cuenta_paginados('PROPIO', q='50300', tipo='cancelados', **filtros)
    assert encontrada['total_resultados'] == 1
    assert encontrada['items'][0]['numero_guia_tauro'] == 50300
    assert cuenta_corriente.movimientos_cuenta_paginados('PROPIO', ambito='nacional')['total_resultados'] == 0
    assert cuenta_corriente.movimientos_cuenta_paginados('PROPIO', hasta='2026-08-31')['total_resultados'] == 0
    libro = load_workbook(BytesIO(generar_excel_cuenta('PROPIO', tipo='cancelados', **filtros)))
    filas = list(libro['Movimientos'].values)[1:]
    assert len(filas) == 2
    assert all(row[7:10] == (0, 0, 0) for row in filas)
    with cuenta_db() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id,estado,monto_ars FROM envios ORDER BY id")
            assert cur.fetchall() == cargos_antes


def _cargo_operativo(db, cliente='DEMO', estado='GUIA_LISTA', nro_fc=''):
    with db() as conn, conn.cursor() as cur:
        cur.execute("INSERT INTO clientes(cliente_id,email) VALUES (%s,'demo@example.invalid') ON CONFLICT DO NOTHING", (cliente,))
        cur.execute("""INSERT INTO solicitudes_guia(cliente_id,estado,producto_alias,destino_pais,
            dest_nombre,dest_direccion,dest_ciudad,dest_zip,remitente_pais,ambito,tracking)
            VALUES(%s,%s,'Muestra','US','Destino','Prueba','Miami','33101','AR','INTERNACIONAL','DEMO') RETURNING id""",(cliente,estado))
        sid=cur.fetchone()['id']
        cur.execute("""INSERT INTO envios(cliente_id,solicitud_id,fecha,monto_ars,estado,ambito,nro_fc)
            VALUES(%s,%s,CURRENT_DATE,100,'ACTIVO','INTERNACIONAL',%s) RETURNING id""",(cliente,sid,nro_fc))
        return sid,cur.fetchone()['id']


def test_anular_admin_retira_cargo_y_estado_juntos_preservando_pago(cuenta_db):
    sid,eid=_cargo_operativo(cuenta_db)
    with cuenta_db() as conn, conn.cursor() as cur:
        cur.execute("INSERT INTO pagos(cliente_id,fecha,monto_ars,estado) VALUES('DEMO',CURRENT_DATE,30,'APROBADO')")
    assert cuenta_corriente.get_facturado_real('DEMO') == 100
    assert cuenta_corriente.cancelar_envio(eid,cliente_id='AJENO') is False
    assert cuenta_corriente.cancelar_envio(eid,cliente_id='DEMO')['estado']=='CANCELADO'
    assert cuenta_corriente.get_facturado_real('DEMO') == 0
    assert cuenta_corriente.total_pagado('DEMO') == 30
    assert cuenta_corriente.cancelar_envio(eid,cliente_id='DEMO') is False
    with cuenta_db() as conn, conn.cursor() as cur:
        cur.execute('SELECT estado FROM solicitudes_guia WHERE id=%s',(sid,))
        assert cur.fetchone()['estado']=='CANCELADO'
        cur.execute('SELECT monto_ars FROM envios WHERE id=%s',(eid,))
        assert cur.fetchone()['monto_ars']==100  # historial intacto
    assert cuenta_corriente.movimientos_cuenta_paginados('DEMO',tipo='cancelados')['total_resultados']==1


def test_baja_admin_se_revierte_si_falla_la_auditoria(cuenta_db,monkeypatch):
    sid,eid=_cargo_operativo(cuenta_db)
    def fallo(*a,**kw):raise RuntimeError('auditoria no disponible')
    monkeypatch.setattr(cuenta_corriente,'registrar_evento_con_cursor',fallo)
    with pytest.raises(RuntimeError):cuenta_corriente.cancelar_envio(eid,cliente_id='DEMO')
    with cuenta_db() as conn,conn.cursor() as cur:
        cur.execute('SELECT s.estado,e.estado AS cargo FROM solicitudes_guia s JOIN envios e ON e.solicitud_id=s.id WHERE s.id=%s',(sid,))
        assert dict(cur.fetchone())=={'estado':'GUIA_LISTA','cargo':'ACTIVO'}


def test_estado_manual_no_puede_dejar_cargo_en_un_envio_cancelado(cuenta_db,monkeypatch):
    from servicios import solicitudes_guia as sg
    monkeypatch.setattr(sg,'get_conn',cuenta_db)
    sid,eid=_cargo_operativo(cuenta_db)
    with pytest.raises(ValueError,match='cargo activo'):sg.actualizar_solicitud_guia(sid,estado='CANCELADO')
    with pytest.raises(ValueError,match='corrección'):sg.actualizar_solicitud_guia(sid,estado='REEMPLAZADO')
    cuenta_corriente.cancelar_envio(eid,cliente_id='DEMO')
    with pytest.raises(ValueError,match='historial'):sg.actualizar_solicitud_guia(sid,estado='GUIA_LISTA',pisar=True)
    assert cuenta_corriente.get_facturado_real('DEMO')==0


def test_facturado_no_se_oculta_al_intentar_cancelarlo(cuenta_db):
    sid,eid=_cargo_operativo(cuenta_db)
    with cuenta_db() as conn,conn.cursor() as cur:
        cur.execute("""INSERT INTO facturas_cliente(cliente_id,tipo,punto_venta,numero,fecha_emision,subtotal,total,pdf,created_by)
            VALUES('DEMO','FC',1,1,CURRENT_DATE,100,100,%s,'qa') RETURNING id""",(b'%PDF-QA',))
        fid=cur.fetchone()['id']
        cur.execute("INSERT INTO facturas_cliente_items(factura_id,envio_id,descripcion,monto) VALUES(%s,%s,'Flete',100)",(fid,eid))
    assert cuenta_corriente.cancelar_envio(eid,cliente_id='DEMO') is False
    assert cuenta_corriente.get_facturado_real('DEMO')==100
    with cuenta_db() as conn,conn.cursor() as cur:
        cur.execute('SELECT estado FROM solicitudes_guia WHERE id=%s',(sid,))
        assert cur.fetchone()['estado']=='GUIA_LISTA'
