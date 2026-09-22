"""Presentación financiera: exactitud, hechos registrados y aislamiento."""

from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal
import os
from urllib.parse import urlsplit
import uuid

import psycopg2
import psycopg2.extras
import pytest

from servicios import experiencia_cuenta as ec


HOY = date(2026, 9, 15)


@pytest.mark.parametrize("vencimiento,pagado,estado,parcial,dias", [
    (date(2026, 9, 10), "20", "VENCIDA", True, 5),
    (HOY, "0", "VENCE_HOY", False, 0),
    (date(2026, 9, 20), "0", "POR_VENCER", False, 0),
    (None, "0", "SIN_FECHA", False, 0),
    (date(2026, 9, 10), "100", "PAGADA", False, 0),
])
def test_vencimiento_no_inventa_fecha_y_distingue_parcial(vencimiento, pagado, estado, parcial, dias):
    f = ec._presentar_factura({
        "id": 42, "punto_venta": 1, "numero": 10, "total": "100",
        "pagado": pagado, "solicitado": "15.25", "fecha_vencimiento": vencimiento,
    }, HOY)
    assert f["estado"] == estado
    assert f["parcial"] is parcial
    assert f["dias_vencida"] == dias
    assert f["saldo_ars"] == Decimal("100") - Decimal(pagado)
    assert f["disponible_pago_ars"] == max(Decimal("0"), f["saldo_ars"] - Decimal("15.25"))
    assert f["destino_pago"] == "F:42"
    if vencimiento is None:
        assert not f["fecha_vencimiento"] and not f["fecha_vencimiento_iso"]


@pytest.mark.parametrize("estado,label", [("PENDIENTE", "En revisión"), ("APROBADO", "Acreditado"), ("RECHAZADO", "Rechazado"), (None, "Acreditado")])
def test_pago_publica_solo_hechos_sin_notas_privadas(estado, label):
    pago = ec._presentar_pago({
        "id": 3, "estado": estado, "fecha": HOY, "monto_ars": "10.10",
        "created_at": datetime(2026, 9, 15, 1, 20, tzinfo=timezone.utc),
        "nota": "Proveedor interno; no publicar", "rechazo_motivo": "Nota administrativa privada",
        "tiene_comprobante": True,
    })
    assert pago["estado_label"] == label
    assert pago["monto_ars"] == Decimal("10.10")
    assert pago["registrado_at"] == "14/09/2026 · 22:20 (AR)"
    assert len(pago["pasos"]) == 2
    assert pago["pasos"][1]["fecha"] == ""
    assert "nota" not in pago and "rechazo_motivo" not in pago
    assert "Proveedor" not in str(pago)


@pytest.mark.parametrize("tope,deuda,reserva,esperado", [
    (None, "100", "20", None), ("-1", "100", "20", None),
    ("0", "0", "0", Decimal("0.00")),
    ("100.10", "60.05", "10.02", Decimal("30.03")),
    ("100", "-50", "10", Decimal("90.00")),
    ("100", "150", "10", Decimal("0.00")),
])
def test_cupo_respeta_tope_cero_credito_y_centavos(tope, deuda, reserva, esperado):
    cupo = ec._presentar_cupo({"tope_deuda_ars": tope, "deuda": deuda, "reservado": reserva})
    assert cupo["disponible_ars"] == esperado
    assert cupo["configurado"] is (esperado is not None)


def test_costos_seis_meses_cruzan_anio_y_ceros_sin_ocultar_ambito():
    costos = ec._presentar_costos([
        {"mes": date(2025, 12, 1), "nacional_ars": "0.10", "internacional_ars": "0.20", "sin_clasificar_ars": "0.30"},
    ], date(2026, 2, 3))
    assert [m["clave"] for m in costos["meses"]] == ["2025-09", "2025-10", "2025-11", "2025-12", "2026-01", "2026-02"]
    assert costos["meses"][-1]["hasta"] == "2026-02-03"
    assert costos["total_ars"] == Decimal("0.60")
    assert costos["hay_sin_clasificar"] is True
    assert costos["meses"][0]["total_ars"] == Decimal("0.00")
    assert costos["periodo_label"] == "Septiembre de 2025–febrero de 2026"


def test_costos_desde_inicio_no_muestran_meses_anteriores_ni_suman_sus_importes():
    costos = ec._presentar_costos([
        {"mes": date(2026, 8, 1), "nacional_ars": "9000"},
        {"mes": date(2026, 9, 1), "nacional_ars": "100", "internacional_ars": "-20"},
    ], HOY, desde=date(2026, 9, 1))
    assert [m["clave"] for m in costos["meses"]] == ["2026-09"]
    assert costos["meses"][0]["desde"] == "2026-09-01"
    assert costos["total_ars"] == Decimal("80.00")
    assert costos["maximo_positivo_ars"] == Decimal("100.00")
    assert costos["maximo_negativo_ars"] == Decimal("20.00")
    assert costos["periodo_label"] == "Desde septiembre de 2026"


def test_costos_inicio_antiguo_conserva_ventana_maxima_de_seis_meses():
    costos = ec._presentar_costos([], date(2027, 5, 10), desde=date(2026, 9, 1))
    assert [m["clave"] for m in costos["meses"]] == ["2026-12", "2027-01", "2027-02", "2027-03", "2027-04", "2027-05"]
    assert costos["periodo_label"] == "Diciembre de 2026–mayo de 2027"


def test_cliente_vacio_se_rechaza_antes_de_abrir_conexion(monkeypatch):
    monkeypatch.setattr(ec, "get_conn", lambda: pytest.fail("No debe abrir DB"))
    with pytest.raises(ValueError):
        ec.obtener_experiencia_cuenta("  ", {})


@pytest.fixture
def cuenta_aislada(monkeypatch):
    url = os.getenv("TAURO_TEST_DATABASE_URL", "").strip()
    if not url:
        pytest.skip("requiere TAURO_TEST_DATABASE_URL local y aislada")
    assert urlsplit(url).hostname in {"127.0.0.1", "localhost", "::1"}, "La prueba sólo acepta PostgreSQL local"
    schema = "test_experiencia_" + uuid.uuid4().hex
    admin = psycopg2.connect(url, cursor_factory=psycopg2.extras.RealDictCursor)
    admin.autocommit = True
    with admin.cursor() as cur:
        cur.execute(f'CREATE SCHEMA "{schema}"')
        cur.execute(f'SET search_path TO "{schema}"')
        # Esquema reducido: permite introducir datos legacy inconsistentes y
        # verificar que el lector filtra propietario/estado por sí mismo.
        cur.execute("""
            CREATE TABLE clientes(cliente_id text PRIMARY KEY, tope_deuda_ars numeric(14,2), puede_emitir boolean);
            CREATE TABLE solicitudes_guia(id integer PRIMARY KEY, cliente_id text, precio_tauro_ars numeric(14,2), tracking text, estado text, cargo_pendiente boolean);
            CREATE TABLE envios(id integer PRIMARY KEY, cliente_id text, fecha date, monto_ars numeric(14,2), estado text, tracking text, solicitud_id integer, ambito text);
            CREATE TABLE ajustes_cliente(id integer PRIMARY KEY, solicitud_id integer, monto_ars numeric(18,4), estado text, aplicado_at timestamptz);
            CREATE TABLE facturas_cliente(id integer PRIMARY KEY, cliente_id text, tipo text, estado text, punto_venta integer, numero integer, fecha_emision date, fecha_vencimiento date, total numeric(14,2));
            CREATE TABLE facturas_cliente_items(id integer PRIMARY KEY, factura_id integer, envio_id integer);
            CREATE TABLE pagos(id integer PRIMARY KEY, cliente_id text, fecha date, created_at timestamptz, monto_ars numeric(14,2), metodo text, referencia text, estado text, comprobante bytea);
            CREATE TABLE pagos_aplicaciones(id integer PRIMARY KEY, pago_id integer, factura_id integer, envio_id integer, monto_ars numeric(14,2), estado text, ambito text);
        """)

    @contextmanager
    def get_conn_aislada():
        conn = psycopg2.connect(url, cursor_factory=psycopg2.extras.RealDictCursor)
        try:
            with conn.cursor() as cur:
                cur.execute(f'SET search_path TO "{schema}"')
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    monkeypatch.setattr(ec, "get_conn", get_conn_aislada)
    monkeypatch.setattr(ec, "_hoy_argentina", lambda: HOY)
    try:
        yield get_conn_aislada
    finally:
        with admin.cursor() as cur:
            cur.execute(f'DROP SCHEMA "{schema}" CASCADE')
        admin.close()


def test_lecturas_reales_aislan_cliente_parciales_reservas_y_costos(cuenta_aislada):
    with cuenta_aislada() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO clientes VALUES ('WAIMAO',1000,true),('OTRO',999999,true);
                INSERT INTO solicitudes_guia VALUES
                    (1,'WAIMAO',100,'W1','GUIA_GENERADA',false),
                    (2,'WAIMAO',200,'W2','GUIA_GENERADA',true),
                    (3,'WAIMAO',50,NULL,'EMITIENDO',false),
                    (4,'WAIMAO',20,NULL,'VERIFICAR_COURIER',false),
                    (5,'WAIMAO',30,'W5','GUIA_GENERADA',true),
                    (6,'OTRO',50000,NULL,'EMITIENDO',false),
                    (7,'WAIMAO',500,'COMPLETA','EMITIENDO',false);
                INSERT INTO envios VALUES
                    (1,'WAIMAO','2026-09-05',100,'ACTIVO','W1',1,'NACIONAL'),
                    (2,'WAIMAO','2026-08-05',200,'ACTIVO','W2',2,'INTERNACIONAL'),
                    (3,'WAIMAO','2026-09-05',999,'CANCELADO','WX',NULL,'NACIONAL'),
                    (4,'WAIMAO','2026-09-05',333,'NC','WN',NULL,'INTERNACIONAL'),
                    (5,'WAIMAO','2026-09-05',70,'ACTIVO','W0',NULL,NULL),
                    (6,'WAIMAO','2026-04-05',10,'ACTIVO','WA',NULL,'NACIONAL'),
                    (7,'WAIMAO','2026-03-05',88,'ACTIVO','WB',NULL,'NACIONAL'),
                    (8,'OTRO','2026-09-05',77777,'ACTIVO','PRIVADO',NULL,'NACIONAL'),
                    (9,'WAIMAO','2026-10-05',25,'ACTIVO','WF',NULL,'NACIONAL');
                INSERT INTO ajustes_cliente VALUES
                    (1,1,20,'APLICADO','2026-09-10 12:00+00'), (2,1,-5,'APLICADO','2026-09-10 12:00+00'),
                    (3,1,999,'PROPUESTO',NULL), (4,2,999,'ANULADO','2026-09-10 12:00+00');
                INSERT INTO facturas_cliente VALUES
                    (1,'WAIMAO','FC','EMITIDA',1,1,'2026-09-01','2026-09-10',100),
                    (2,'WAIMAO','FC','EMITIDA',1,2,'2026-09-01','2026-09-20',200),
                    (3,'WAIMAO','FC','EMITIDA',1,3,'2026-09-01','2026-09-01',100),
                    (4,'WAIMAO','NC','EMITIDA',1,4,'2026-09-01','2026-09-01',500),
                    (5,'WAIMAO','FC','ANULADA',1,5,'2026-09-01','2026-09-01',500),
                    (6,'WAIMAO','FC','EMITIDA',1,6,'2026-09-01',NULL,50),
                    (7,'OTRO','FC','EMITIDA',1,7,'2026-09-01','2026-09-01',50000);
                INSERT INTO facturas_cliente_items VALUES (1,1,1),(2,1,1),(3,2,2);
                INSERT INTO pagos VALUES
                    (1,'WAIMAO','2026-09-11','2026-09-11 12:00+00',40,'transferencia','R1','APROBADO',NULL),
                    (2,'WAIMAO','2026-09-12','2026-09-12 12:00+00',30,'transferencia','R2','APROBADO',NULL),
                    (3,'WAIMAO','2026-09-13','2026-09-13 12:00+00',20,'transferencia','R3','PENDIENTE',NULL),
                    (4,'WAIMAO','2026-09-14','2026-09-14 12:00+00',500,'transferencia','R4','RECHAZADO',NULL),
                    (5,'WAIMAO','2026-09-10','2026-09-10 12:00+00',100,'transferencia','R5','APROBADO',NULL),
                    (6,'WAIMAO','2026-09-09','2026-09-09 12:00+00',15,'transferencia','R6',NULL,NULL),
                    (7,'OTRO','2026-09-15','2026-09-15 12:00+00',1000,'transferencia','PRIVADO','APROBADO',NULL);
                INSERT INTO pagos_aplicaciones VALUES
                    (1,1,1,NULL,40,'APLICADA','NACIONAL'),
                    (2,2,NULL,1,30,'APLICADA','NACIONAL'),
                    (3,3,1,NULL,20,'SOLICITADA','NACIONAL'),
                    (4,4,1,NULL,500,'APLICADA','NACIONAL'),
                    (5,5,3,NULL,100,'APLICADA','NACIONAL'),
                    (6,6,2,NULL,15,'APLICADA','INTERNACIONAL'),
                    (7,7,1,NULL,1000,'APLICADA','NACIONAL'),
                    (8,1,7,NULL,5,'APLICADA','NACIONAL'),
                    (9,1,NULL,8,5,'APLICADA','NACIONAL');
            """)
    resultado = ec.obtener_experiencia_cuenta(" waimao ", {"pagos_pendientes_ars": Decimal("20")})
    v = resultado["vencimientos"]
    assert [f["id"] for f in v["items"]] == [1, 2, 6]
    assert v["cantidad"] == 3
    assert v["total_pendiente_ars"] == Decimal("265.00")
    assert v["total_vencido_ars"] == Decimal("30.00")
    assert v["total_en_revision_ars"] == Decimal("20.00")
    assert v["items"][0]["pagado_ars"] == Decimal("70.00")
    assert v["items"][0]["disponible_pago_ars"] == Decimal("10.00")
    assert resultado["pagos"][0]["estado"] == "RECHAZADO"
    assert resultado["pagos"][0]["aplicaciones"] == []
    assert resultado["pagos"][-1]["estado"] == "APROBADO"
    assert "PRIVADO" not in str(resultado)
    p1 = next(p for p in resultado["pagos"] if p["id"] == 1)
    assert len(p1["aplicaciones"]) == 1  # No factura/envío de OTRO.
    # El cupo incluye los ajustes aplicados (+20 -5), igual que el libro.
    assert resultado["cupo"]["deuda_ars"] == Decimal("323.00")
    assert resultado["cupo"]["reservado_ars"] == Decimal("100.00")
    assert resultado["cupo"]["disponible_ars"] == Decimal("577.00")
    costos = resultado["costos"]
    assert costos["total_ars"] == Decimal("185.00")  # Septiembre: 100+20-5+70.
    assert [m["clave"] for m in costos["meses"]] == ["2026-09"]
    assert costos["meses"][-1]["nacional_ars"] == Decimal("115.00")
    assert costos["meses"][-1]["sin_clasificar_ars"] == Decimal("70.00")
    assert ec.obtener_experiencia_cuenta("WAIMAO' OR '1'='1", {})["pagos"] == []


def test_totales_vencimientos_no_se_recortan_al_limitar_seis(cuenta_aislada):
    with cuenta_aislada() as conn:
        with conn.cursor() as cur:
            cur.execute("INSERT INTO clientes VALUES ('WAIMAO',NULL,false)")
            cur.execute("""
                INSERT INTO facturas_cliente
                SELECT n,'WAIMAO','FC','EMITIDA',1,n,'2026-09-01','2026-09-10',10.01
                FROM generate_series(1,9) AS n
            """)
    resultado = ec.obtener_experiencia_cuenta("WAIMAO", {})
    assert len(resultado["vencimientos"]["items"]) == 6
    assert resultado["vencimientos"]["cantidad"] == 9
    assert resultado["vencimientos"]["cantidad_vencida"] == 9
    assert resultado["vencimientos"]["total_vencido_ars"] == Decimal("90.09")
    assert not resultado["cupo"]["configurado"]
    assert len(resultado["costos"]["meses"]) == 1


def test_inicio_waimao_filtra_pagos_por_fecha_antes_del_limite_sin_cambiar_deuda(cuenta_aislada):
    with cuenta_aislada() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO clientes VALUES ('WAIMAO',500,true),('OTRO',500,true);
                INSERT INTO envios VALUES
                    (1,'WAIMAO','2026-08-31',100,'ACTIVO','ANTERIOR',NULL,'NACIONAL'),
                    (2,'WAIMAO','2026-09-01',200,'ACTIVO','ACTUAL',NULL,'NACIONAL'),
                    (3,'OTRO','2026-08-31',999,'ACTIVO','OTRO-ANTERIOR',NULL,'NACIONAL');
                INSERT INTO facturas_cliente VALUES
                    (1,'WAIMAO','FC','EMITIDA',1,1,'2026-08-30','2026-08-31',100);
                INSERT INTO pagos VALUES
                    (1,'WAIMAO','2026-08-31','2026-09-14 12:00+00',20,'Transferencia','ANTERIOR','APROBADO',NULL),
                    (2,'WAIMAO','2026-09-01','2026-08-31 12:00+00',30,'Transferencia','INICIO','APROBADO',NULL),
                    (3,'WAIMAO','2026-08-31','2026-09-15 12:00+00',10,'Transferencia','PENDIENTE-ANTERIOR','PENDIENTE',NULL),
                    (4,'WAIMAO','2026-09-01','2026-09-01 12:00+00',40,'Transferencia','RECHAZADO-INICIO','RECHAZADO',NULL),
                    (5,'OTRO','2026-08-31','2026-09-15 12:00+00',30,'Transferencia','OTRO','APROBADO',NULL),
                    (6,'WAIMAO','2026-09-02','2026-09-02 12:00+00',15,'Transferencia','POSTERIOR','APROBADO',NULL);
                INSERT INTO pagos
                SELECT n,'WAIMAO','2026-08-31','2026-09-15 13:00+00',1,
                       'Transferencia','LEGACY-'||n,'RECHAZADO',NULL
                FROM generate_series(10,18) AS n;
            """)
    resultado = ec.obtener_experiencia_cuenta(" waimao ", {"pagos_pendientes_ars": Decimal("10")})
    assert {p["id"] for p in resultado["pagos"]} == {2, 4, 6}
    assert {p["fecha"] for p in resultado["pagos"]} == {"01/09/2026", "02/09/2026"}
    assert resultado["pagos_en_revision_ars"] == Decimal("10.00")
    assert resultado["vencimientos"]["items"][0]["fecha_vencimiento"] == "31/08/2026"
    assert resultado["cupo"]["deuda_ars"] == Decimal("235.00")
    assert resultado["cupo"]["disponible_ars"] == Decimal("265.00")
    assert resultado["costos"]["total_ars"] == Decimal("200.00")
    otro = ec.obtener_experiencia_cuenta("OTRO", {})
    assert [p["id"] for p in otro["pagos"]] == [5]
    assert len(otro["costos"]["meses"]) == 6
    assert otro["costos"]["total_ars"] == Decimal("999.00")
    assert otro["costos"]["periodo_label"] == "Abril–septiembre de 2026"
    with cuenta_aislada() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS cantidad FROM pagos WHERE cliente_id='WAIMAO'")
            assert cur.fetchone()["cantidad"] == 14


@pytest.fixture
def movimientos_aislados(cuenta_aislada, monkeypatch):
    from servicios import cuenta_corriente as cc
    with cuenta_aislada() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                ALTER TABLE envios ADD COLUMN created_at timestamptz DEFAULT NOW(),
                    ADD COLUMN nro_fc text, ADD COLUMN descripcion text,
                    ADD COLUMN factura_pdf bytea;
                ALTER TABLE solicitudes_guia ADD COLUMN visible_cliente boolean DEFAULT true,
                    ADD COLUMN test boolean DEFAULT false, ADD COLUMN dest_nombre text,
                    ADD COLUMN remitente_nombre text, ADD COLUMN etiqueta_cliente text,
                    ADD COLUMN remitente_ciudad text, ADD COLUMN remitente_pais text,
                    ADD COLUMN dest_ciudad text, ADD COLUMN destino_pais text,
                    ADD COLUMN created_at timestamptz DEFAULT NOW(), ADD COLUMN ambito text,
                    ADD COLUMN numero_guia_tauro bigint, ADD COLUMN producto_alias text,
                    ADD COLUMN cantidad integer DEFAULT 1;
                ALTER TABLE facturas_cliente ADD COLUMN pdf bytea;
                ALTER TABLE facturas_cliente_items ADD COLUMN ajuste_id integer;
                ALTER TABLE pagos_aplicaciones ADD COLUMN updated_at timestamptz DEFAULT NOW();
                ALTER TABLE ajustes_cliente ADD COLUMN tipo text DEFAULT 'DEBITO',
                    ADD COLUMN conciliacion_id integer, ADD COLUMN motivo text,
                    ADD COLUMN origen text NOT NULL DEFAULT 'CONCILIACION_COURIER',
                    ADD COLUMN precio_anterior_ars numeric(18,4) DEFAULT 100,
                    ADD COLUMN precio_nuevo_ars numeric(18,4) DEFAULT 100;
                CREATE TABLE conciliaciones_envio(
                    id integer PRIMARY KEY, tax_cliente_ars numeric DEFAULT 0,
                    diferencia_flete_ars numeric DEFAULT 0, motivo_diferencia text DEFAULT 'PESO_REAL',
                    precio_cliente_inicial_ars numeric DEFAULT 100,
                    peso_cotizado_kg numeric, peso_final_facturado_kg numeric, peso_base_facturado text);
                CREATE TABLE factura_courier_item_matches(solicitud_id integer,item_id integer,estado text);
                CREATE TABLE facturas_courier_items(id integer,descripcion text,concepto_tipo text);
            """)
    monkeypatch.setattr(cc, "get_conn", cuenta_aislada)
    return cuenta_aislada


def test_buscar_pago_por_factura_guia_o_destinatario_no_duplica_haber(movimientos_aislados):
    from servicios import cuenta_corriente as cc
    from servicios.export_cuenta import generar_excel_cuenta
    from io import BytesIO
    from openpyxl import load_workbook
    with movimientos_aislados() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO clientes VALUES ('WAIMAO',1000,true),('OTRO',1000,true);
                INSERT INTO solicitudes_guia(id,cliente_id,tracking,dest_nombre)
                    VALUES (1,'WAIMAO','GUIA-W1','DESTINATARIO_W1'),
                           (2,'WAIMAO','GUIA-W2','DESTINATARIO_W2'),
                           (3,'OTRO','PRIVADO-AJENO','PRIVADO-AJENO');
                INSERT INTO envios(id,cliente_id,fecha,monto_ars,estado,solicitud_id,ambito)
                    VALUES (1,'WAIMAO','2026-09-10',100,'ACTIVO',1,'NACIONAL'),
                           (2,'WAIMAO','2026-09-10',100,'ACTIVO',2,'NACIONAL'),
                           (3,'OTRO','2026-09-10',100,'ACTIVO',3,'NACIONAL');
                INSERT INTO facturas_cliente(id,cliente_id,tipo,estado,punto_venta,numero,total)
                    VALUES (1,'WAIMAO','FC','EMITIDA',1,42,100),
                           (2,'WAIMAO','FC','EMITIDA',1,43,100),
                           (3,'OTRO','FC','EMITIDA',1,99999,100);
                INSERT INTO facturas_cliente_items(id,factura_id,envio_id)
                    VALUES (1,1,1),(2,1,1),(3,2,2),(4,3,3);
                INSERT INTO pagos(id,cliente_id,fecha,created_at,monto_ars,metodo,estado)
                    VALUES (1,'WAIMAO','2026-09-14','2026-09-14 15:00+00',100,'Transferencia','APROBADO'),
                           (2,'WAIMAO','2026-09-14','2026-09-14 15:00+00',20,'Transferencia','PENDIENTE'),
                           (3,'WAIMAO','2026-09-14','2026-09-14 15:00+00',100,'Transferencia','RECHAZADO'),
                           (4,'OTRO','2026-09-14','2026-09-14 15:00+00',100,'Transferencia','APROBADO'),
                           (5,'WAIMAO','2026-09-14','2026-09-14 15:00+00',30,'Transferencia',NULL);
                INSERT INTO pagos_aplicaciones(id,pago_id,factura_id,envio_id,monto_ars,estado,ambito)
                    VALUES (1,1,1,NULL,40,'APLICADA','NACIONAL'),
                           (2,1,2,NULL,60,'APLICADA','NACIONAL'),
                           (3,2,1,NULL,15,'SOLICITADA','NACIONAL'),
                           (4,3,1,NULL,100,'APLICADA','NACIONAL'),
                           (5,4,1,NULL,100,'APLICADA','NACIONAL'),
                           (6,1,3,NULL,5,'APLICADA','NACIONAL'),
                           (7,5,NULL,2,30,'APLICADA','NACIONAL');
            """)
    for consulta in ("FC 0001-00000042", "GUIA-W1", "DESTINATARIO_W1"):
        r = cc.movimientos_cuenta_paginados("WAIMAO", tipo="pagos", q=consulta)
        assert r["total_resultados"] == 2
        assert {m["pago_id"] for m in r["items"]} == {1, 2}
        assert sum(m["haber_ars"] for m in r["items"]) == Decimal("40.00")
        assert sum(m["debe_ars"] for m in r["items"]) == 0
    r = cc.movimientos_cuenta_paginados("WAIMAO", tipo="pagos", q="GUIA-W2")
    assert r["total_resultados"] == 2
    assert sum(m["haber_ars"] for m in r["items"]) == Decimal("90.00")
    for consulta in ("PRIVADO-AJENO", "00099999"):
        assert cc.movimientos_cuenta_paginados("WAIMAO", tipo="pagos", q=consulta)["total_resultados"] == 0
    libro = load_workbook(BytesIO(generar_excel_cuenta("WAIMAO", tipo="pagos", q="GUIA-W1")))
    filas = list(libro["Movimientos"].values)[1:]
    assert len(filas) == 2
    assert sum(f[8] for f in filas) == 40


def test_costos_asientos_negativos_y_timezone_coinciden_con_lista_y_excel(movimientos_aislados):
    from servicios import cuenta_corriente as cc
    from servicios.export_cuenta import generar_excel_cuenta
    from io import BytesIO
    from openpyxl import load_workbook
    with movimientos_aislados() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO clientes VALUES ('CLIENTE_QA',1000,true);
                INSERT INTO solicitudes_guia(id,cliente_id,tracking)
                    VALUES (1,'CLIENTE_QA','ANTERIOR'),(2,'CLIENTE_QA','ACTUAL');
                INSERT INTO envios(id,cliente_id,fecha,monto_ars,estado,solicitud_id,ambito)
                    VALUES (1,'CLIENTE_QA','2026-03-10',100,'ACTIVO',1,'NACIONAL'),
                           (2,'CLIENTE_QA','2026-09-10',50,'ACTIVO',2,'INTERNACIONAL');
                INSERT INTO conciliaciones_envio(id) SELECT generate_series(1,5);
                INSERT INTO ajustes_cliente(id,solicitud_id,monto_ars,estado,aplicado_at,tipo,conciliacion_id)
                    VALUES (1,1,20,'APLICADO','2026-09-15 01:00+00','DEBITO',1),
                           (2,1,-75,'APLICADO','2026-09-15 01:00+00','CREDITO',2),
                           (3,1,-7,'APLICADO','2026-08-15 12:00+00','CREDITO',3),
                           (4,1,0.005,'APLICADO','2026-09-15 01:00+00','DEBITO',4),
                           (5,1,0.005,'APLICADO','2026-09-15 01:00+00','DEBITO',5);
            """)
    costos = ec.obtener_experiencia_cuenta("CLIENTE_QA", {})["costos"]
    agosto, septiembre = costos["meses"][-2:]
    assert agosto["total_ars"] == Decimal("-7.00")
    assert septiembre["nacional_ars"] == Decimal("-54.98")
    assert septiembre["internacional_ars"] == Decimal("50.00")
    assert septiembre["total_ars"] == Decimal("-4.98")
    assert septiembre["positivo_ars"] == Decimal("50.00")
    assert septiembre["negativo_ars"] == Decimal("54.98")
    assert costos["total_ars"] == Decimal("-11.98")
    assert costos["maximo_positivo_ars"] == Decimal("50.00")
    assert costos["maximo_negativo_ars"] == Decimal("54.98")
    assert costos["hay_negativos"] is True
    for mes in (agosto, septiembre):
        filtros = {"tipo": "costos", "desde": mes["desde"], "hasta": mes["hasta"]}
        r = cc.movimientos_cuenta_paginados("CLIENTE_QA", **filtros)
        assert sum(m["debe_ars"]-m["haber_ars"] for m in r["items"]) == mes["total_ars"]
        libro = load_workbook(BytesIO(generar_excel_cuenta("CLIENTE_QA", **filtros)))
        filas = list(libro["Movimientos"].values)[1:]
        assert sum(Decimal(str(f[7]))-Decimal(str(f[8])) for f in filas) == mes["total_ars"]
    ajustes = cc.movimientos_cuenta_paginados("CLIENTE_QA", tipo="diferencias", desde="2026-09-14", hasta="2026-09-14")
    assert ajustes["total_resultados"] == 4
    assert {m["fecha_iso"] for m in ajustes["items"]} == {"2026-09-14"}
    assert cc.movimientos_cuenta_paginados("CLIENTE_QA", tipo="diferencias", desde="2026-09-15")["total_resultados"] == 0
